from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def test_excel_report_exists():
    files = list((ROOT / "outputs").glob("fund_performance_report_*.xlsx"))
    assert files, "Excel report missing"

def test_workbook_has_expected_sheets():
    report = max(
        (ROOT / "outputs").glob("fund_performance_report_*.xlsx"),
        key=lambda p: p.stat().st_mtime
    )
    wb = load_workbook(report, read_only=True)
    assert {"Executive Summary", "Monthly Performance", "Daily Returns",
            "Data Quality", "Validated Sample"}.issubset(wb.sheetnames)
    wb.close()

def test_summary_has_performance_metrics():
    report = max((ROOT / "outputs").glob("fund_performance_report_*.xlsx"),
                 key=lambda p: p.stat().st_mtime)
    df = pd.read_excel(report, sheet_name="Executive Summary")
    assert {"Compounded proxy return", "Annualized volatility",
            "Maximum drawdown"}.issubset(set(df["metric"]))

def test_audit_log_exists():
    log = ROOT / "logs" / "audit_history.jsonl"
    assert log.exists() and log.stat().st_size > 0

def test_daily_returns_are_present():
    report = max((ROOT / "outputs").glob("fund_performance_report_*.xlsx"),
                 key=lambda p: p.stat().st_mtime)
    df = pd.read_excel(report, sheet_name="Daily Returns")
    assert not df.empty
    assert df["portfolio_return"].notna().all()

from openpyxl import Workbook, load_workbook

from src.remediation import create_issue, list_issues, update_issue
from src.report_with_remediation import refresh_remediation_sheet


def test_tracker_uses_persistent_records_and_preserves_other_sheets(tmp_path):
    db = tmp_path / "remediation.db"
    workbook_path = tmp_path / "report.xlsx"

    create_issue(
        "DQ-MANUAL-001",
        "Manual control exception",
        severity="HIGH",
        owner="Reporting Team",
        db_path=db,
    )
    update_issue(
        "DQ-MANUAL-001",
        status="IN_PROGRESS",
        actor="analyst",
        details="Investigation started",
        db_path=db,
    )

    workbook = Workbook()
    data_quality = workbook.active
    data_quality.title = "Data Quality"
    data_quality.append(["exception_type", "details"])
    data_quality.append(["missing_value", "A source value is missing"])
    workbook.create_sheet("Executive Summary")
    workbook.save(workbook_path)

    count = refresh_remediation_sheet(workbook_path, db_path=db)

    result = load_workbook(workbook_path, read_only=True, data_only=True)
    assert "Executive Summary" in result.sheetnames
    assert "Remediation Tracker" in result.sheetnames

    rows = list(
        result["Remediation Tracker"].iter_rows(values_only=True)
    )
    headers = list(rows[0])
    records = [dict(zip(headers, row)) for row in rows[1:]]

    manual = next(
        item for item in records
        if item["issue_key"] == "DQ-MANUAL-001"
    )
    assert manual["status"] == "IN_PROGRESS"
    assert manual["owner"] == "Reporting Team"
    assert count == 2
    assert len(list_issues(db_path=db)) == 2
    result.close()


def test_refresh_does_not_duplicate_existing_exceptions(tmp_path):
    db = tmp_path / "remediation.db"
    workbook_path = tmp_path / "report.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data Quality"
    sheet.append(["exception_type", "details"])
    sheet.append(["duplicate_date", "Duplicate ticker/date record"])
    workbook.save(workbook_path)

    first_count = refresh_remediation_sheet(workbook_path, db_path=db)
    first_issue_count = len(list_issues(db_path=db))

    second_count = refresh_remediation_sheet(workbook_path, db_path=db)
    second_issue_count = len(list_issues(db_path=db))

    assert first_count == second_count == 1
    assert first_issue_count == second_issue_count == 1

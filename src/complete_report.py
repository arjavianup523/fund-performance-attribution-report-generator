from pathlib import Path
from datetime import datetime, timezone
import json
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "raw"
OUT = ROOT / "outputs"
LOG = ROOT / "logs"

def calculate_attribution(path):
    d = pd.read_csv(path)
    required = {"sector","portfolio_weight","benchmark_weight","portfolio_return","benchmark_return"}
    if not required.issubset(d.columns):
        raise ValueError(f"Attribution input missing columns: {required-set(d.columns)}")
    numeric = ["portfolio_weight","benchmark_weight","portfolio_return","benchmark_return"]
    for c in numeric:
        d[c] = pd.to_numeric(d[c], errors="raise")
    if d[numeric].isna().any().any():
        raise ValueError("Attribution inputs contain missing numeric values.")
    if (d[["portfolio_weight","benchmark_weight"]] < 0).any().any():
        raise ValueError("Weights cannot be negative.")
    if not np.isclose(d["portfolio_weight"].sum(), 1, atol=1e-6):
        raise ValueError("Portfolio weights must sum to 1.")
    if not np.isclose(d["benchmark_weight"].sum(), 1, atol=1e-6):
        raise ValueError("Benchmark weights must sum to 1.")
    if d["sector"].duplicated().any():
        raise ValueError("Each sector must appear only once.")

    benchmark_total = (d.benchmark_weight * d.benchmark_return).sum()
    portfolio_total = (d.portfolio_weight * d.portfolio_return).sum()
    d["allocation_effect"] = (d.portfolio_weight-d.benchmark_weight) * (d.benchmark_return-benchmark_total)
    d["selection_effect"] = d.benchmark_weight * (d.portfolio_return-d.benchmark_return)
    d["interaction_effect"] = (d.portfolio_weight-d.benchmark_weight) * (d.portfolio_return-d.benchmark_return)
    d["total_effect"] = d[["allocation_effect","selection_effect","interaction_effect"]].sum(axis=1)
    d["portfolio_contribution"] = d.portfolio_weight*d.portfolio_return
    d["benchmark_contribution"] = d.benchmark_weight*d.benchmark_return

    totals = pd.DataFrame([
        ["Portfolio return", portfolio_total],
        ["Benchmark return", benchmark_total],
        ["Active return", portfolio_total-benchmark_total],
        ["Total allocation effect", d.allocation_effect.sum()],
        ["Total selection effect", d.selection_effect.sum()],
        ["Total interaction effect", d.interaction_effect.sum()],
        ["Attribution reconciliation difference", (portfolio_total-benchmark_total)-d.total_effect.sum()]
    ], columns=["metric","value"])
    if abs(float(totals.loc[totals.metric=="Attribution reconciliation difference","value"].iloc[0])) > 1e-8:
        raise ArithmeticError("Attribution does not reconcile to active return.")
    return d, totals

def main():
    OUT.mkdir(exist_ok=True)
    LOG.mkdir(exist_ok=True)
    input_file = DATA / "all_stocks_5yr.csv"
    raw = pd.read_csv(input_file)
    required = {"date","open","high","low","close","volume","Name"}
    if not required.issubset(raw.columns):
        raise ValueError(f"Stock file missing columns: {required-set(raw.columns)}")

    raw["source_row"] = np.arange(2,len(raw)+2)
    raw["date"] = pd.to_datetime(raw.date, errors="coerce", dayfirst=True)
    raw["Name"] = raw.Name.astype("string").str.strip().str.upper()
    for c in ["open","high","low","close","volume"]:
        raw[c] = pd.to_numeric(raw[c],errors="coerce")

    exceptions = []
    for c in ["date","Name","close"]:
        bad = raw[c].isna() | (raw[c].eq("") if c=="Name" else False)
        for _, r in raw.loc[bad].iterrows():
            exceptions.append({"source_row":int(r.source_row),"rule":f"INVALID_{c.upper()}","severity":"CRITICAL","status":"OPEN","detail":f"Missing or invalid {c}"})
    bad_price = raw.close.notna() & (raw.close<=0)
    for _,r in raw.loc[bad_price].iterrows():
        exceptions.append({"source_row":int(r.source_row),"rule":"NONPOSITIVE_CLOSE","severity":"CRITICAL","status":"OPEN","detail":"Close price must be positive"})
    dup = raw.duplicated(["Name","date"],keep="first") & raw.Name.notna() & raw.date.notna()
    for _,r in raw.loc[dup].iterrows():
        exceptions.append({"source_row":int(r.source_row),"rule":"DUPLICATE_TICKER_DATE","severity":"HIGH","status":"OPEN","detail":"Duplicate ticker/date excluded"})

    valid = raw.loc[raw.date.notna() & raw.Name.notna() & raw.Name.ne("") & raw.close.notna() & (raw.close>0) & ~dup].copy()
    valid = valid.sort_values(["Name","date"])
    valid["daily_return"] = valid.groupby("Name").close.pct_change()
    returns = valid.dropna(subset=["daily_return"])
    daily = returns.groupby("date",as_index=False).daily_return.mean().rename(columns={"daily_return":"portfolio_return"})
    daily = daily.sort_values("date")
    daily["growth_of_1"] = (1+daily.portfolio_return).cumprod()
    daily["drawdown"] = daily.growth_of_1/daily.growth_of_1.cummax()-1
    daily["month"] = daily.date.dt.to_period("M").astype(str)
    monthly = daily.groupby("month",as_index=False).agg(
        portfolio_return=("portfolio_return",lambda x:(1+x).prod()-1),
        observations=("portfolio_return","count"),
        start_date=("date","min"),end_date=("date","max"))

    r = daily.portfolio_return
    total = (1+r).prod()-1
    vol = r.std(ddof=1)*np.sqrt(252)
    ann = (1+total)**(252/len(r))-1 if len(r) and total>-1 else np.nan
    sharpe = r.mean()/r.std(ddof=1)*np.sqrt(252) if len(r)>1 and r.std(ddof=1)>0 else np.nan
    stock_summary = pd.DataFrame([
        ["Scope","Equal-weighted stock-universe proxy, NOT a mutual fund"],
        ["Valid price rows",len(valid)],["Unique tickers",valid.Name.nunique()],
        ["Start date",valid.date.min()],["End date",valid.date.max()],
        ["Compounded proxy return",total],["Annualized return (252-day approximation)",ann],
        ["Annualized volatility",vol],["Maximum drawdown",daily.drawdown.min()],
        ["Sharpe ratio (risk-free rate assumed zero)",sharpe],
        ["Benchmark comparison","Not available for the stock universe in this input"],
        ["Data quality exceptions",len(exceptions)]
    ],columns=["metric","value"])

    attribution, attribution_summary = calculate_attribution(DATA/"illustrative_attribution_inputs.csv")
    attribution.insert(0,"data_status","ILLUSTRATIVE SAMPLE — NOT ACTUAL CLIENT/FUND DATA")
    attribution_summary.insert(0,"data_status","ILLUSTRATIVE SAMPLE — NOT ACTUAL CLIENT/FUND DATA")
    exception_df = pd.DataFrame(exceptions,columns=["source_row","rule","severity","status","detail"])
    if exception_df.empty:
        exception_df = pd.DataFrame([{"source_row":"","rule":"NO_EXCEPTIONS","severity":"INFO","status":"CLOSED","detail":"No exceptions found"}])
    remediation = exception_df.copy()
    remediation["owner"] = "UNASSIGNED"
    remediation["action_required"] = "Investigate source record and document resolution"
    remediation["resolution_notes"] = ""

    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    report = OUT/f"fund_performance_report_{run_id}.xlsx"
    with pd.ExcelWriter(report,engine="openpyxl") as writer:
        stock_summary.to_excel(writer,sheet_name="Executive Summary",index=False)
        monthly.to_excel(writer,sheet_name="Monthly Performance",index=False)
        daily.to_excel(writer,sheet_name="Daily Returns",index=False)
        exception_df.to_excel(writer,sheet_name="Data Quality",index=False)
        remediation.to_excel(writer,sheet_name="Remediation Tracker",index=False)
        attribution.to_excel(writer,sheet_name="Sector Attribution",index=False)
        attribution_summary.to_excel(writer,sheet_name="Attribution Summary",index=False)
        valid.head(1000).to_excel(writer,sheet_name="Validated Sample",index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=Font(bold=True,color="FFFFFF")
                cell.fill=PatternFill("solid",fgColor="17365D")
                cell.alignment=Alignment(wrap_text=True)
            for col in ws.columns:
                letter=get_column_letter(col[0].column)
                ws.column_dimensions[letter].width=min(max(max(len(str(c.value or "")) for c in col)+2,12),48)

    audit={"run_id":run_id,"timestamp_utc":datetime.now(timezone.utc).isoformat(),
           "input_rows":len(raw),"valid_rows":len(valid),"ticker_count":int(valid.Name.nunique()),
           "exception_count":len(exceptions),"report":str(report),
           "attribution_source":"illustrative_attribution_inputs.csv"}
    with (LOG/"audit_history.jsonl").open("a",encoding="utf-8") as f:
        f.write(json.dumps(audit,default=str)+"\n")
    print("REPORT:",report)
    print("VALID ROWS:",len(valid),"| TICKERS:",valid.Name.nunique(),"| EXCEPTIONS:",len(exceptions))
    print("ATTRIBUTION RECONCILIATION:")
    print(attribution_summary.to_string(index=False))
    print("AUDIT:",LOG/"audit_history.jsonl")

if __name__=="__main__":
    main()
